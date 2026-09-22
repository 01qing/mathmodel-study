"""Read back exported workbooks; do not modify them."""
from pathlib import Path
import json, hashlib
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
reports = []
for key in ['result11', 'result12', 'result21', 'result22']:
    matrix = json.loads((ROOT/'results'/f'{key}_matrix.json').read_text(encoding='utf-8'))
    path = ROOT/'deliverables'/f'{key}.xlsx'
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    assert len(book.sheetnames) == 1
    sheet = book.active
    sheet.calculate_dimension(force=True)
    assert (sheet.max_row, sheet.max_column) == (len(matrix), len(matrix[0]))
    count = 0
    for i, row in enumerate(sheet.iter_rows(min_row=1, max_row=len(matrix), min_col=1, max_col=len(matrix[0]))):
        for j, cell in enumerate(row):
            assert cell.data_type not in ('f', 'e'), (key, i, j, cell.value)
            assert cell.value == matrix[i][j], (key, i, j, cell.value, matrix[i][j])
            count += 1
    book.close()
    reports.append(dict(id=key, status='PASS', cells_compared=count, rows=len(matrix), columns=len(matrix[0]), sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    print(key, 'PASS', count, flush=True)
(ROOT/'results'/'XLSX_VALIDATION.json').write_text(json.dumps(reports, indent=2), encoding='utf-8')
