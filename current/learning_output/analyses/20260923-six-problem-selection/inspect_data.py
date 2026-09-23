from pathlib import Path
import json,zipfile,io,csv,collections,hashlib,sys
import logging
logging.getLogger('pypdf').setLevel(logging.ERROR)
sys.path.insert(0,str(Path(__file__).parent/'runtime_deps'))
import numpy as np
from scipy.io import whosmat,loadmat
from openpyxl import load_workbook
from pypdf import PdfReader
ROOT=Path(r'D:\数模\数模练习'); OUT=Path(__file__).parent
sys.stdout.reconfigure(encoding='utf-8')
report={}
for letter in 'DF':
    for p in (ROOT/f'{letter}题exp').rglob('*.pdf'):
        r=PdfReader(p)
        target=OUT/f'{letter}_{p.stem}_PDF提取.txt'
        if not target.exists():
            text='\n\n'.join(f'=== PAGE {i+1} ===\n{page.extract_text()}' for i,page in enumerate(r.pages))
            target.write_text(text,encoding='utf-8')
        report.setdefault('pdfs',[]).append({'path':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'pages':len(r.pages),'output':target.name})
def excel_info(src):
    wb=load_workbook(src,read_only=True,data_only=True)
    out=[]
    for s in wb:
        rows=[list(r) for r in s.iter_rows(values_only=True)]
        rows=[r for r in rows if any(x is not None for x in r)]
        out.append({'sheet':s.title,'rows':s.max_row,'cols':s.max_column,'nonempty_rows':len(rows),'sample':rows[:8]})
    wb.close();return out
for p in (ROOT/'D题exp').rglob('*.xlsx'):
    report.setdefault('D_excel',{})[str(p.relative_to(ROOT/'D题exp'))]=excel_info(p)
for letter in 'AB':
    zp=next((ROOT/f'{letter}题exp').glob('*.zip'))
    with zipfile.ZipFile(zp) as z:
        if letter=='B':
            report['B_excel']={n:excel_info(io.BytesIO(z.read(n))) for n in z.namelist() if n.endswith('.xlsx')}
        else:
            report['A_graphs']=[]
            for n in z.namelist():
                if n.startswith('data/') and n.endswith('.json'):
                    d=json.loads(z.read(n)); report['A_graphs'].append({'name':n,'sizes':{k:len(v) for k,v in d.items() if isinstance(v,(dict,list))}})
            for n in ['README.md','data/config.txt']:
                (OUT/f'A_{Path(n).name}').write_bytes(z.read(n))
            d=json.loads(z.read('data/case_001.json'))
            report['A_sample']={k:v[:2] if isinstance(v,list) else v for k,v in d.items()}
for p in (ROOT/'C题exp').glob('*.mat'):
    obj={'vars':whosmat(p)}
    d=loadmat(p)
    obj['values']={}
    for k,v in d.items():
        if not k.startswith('__'):
            item={'shape':v.shape,'dtype':str(v.dtype)}
            if v.dtype.names:item['fields']=v.dtype.names
            if v.dtype.kind in 'if' and v.size:
                item['finite_fraction']=float(np.isfinite(v).mean())
                if v.ndim==2 and min(v.shape)==10:
                    arr=v if v.shape[1]==10 else v.T
                    item['visual_events']=dict(zip(*[x.tolist() for x in np.unique(arr[:,7],return_counts=True)]))
                    item['response_events']=dict(zip(*[x.tolist() for x in np.unique(arr[:,8],return_counts=True)]))
            obj['values'][k]=item
    report.setdefault('C_mat',{})[p.name]=obj
for p in (ROOT/'D题exp').rglob('*.mat'):
    try: report.setdefault('D_mat',{})[p.name]=whosmat(p)
    except Exception as e: report.setdefault('D_mat',{})[p.name]={'error':str(e)}
try:
    import rasterio
    p=next((ROOT/'D题exp').rglob('*.tif'))
    with rasterio.open(p) as ds:
        report['D_DEM']={'shape':ds.shape,'crs':str(ds.crs),'bounds':list(ds.bounds),'resolution':ds.res,'nodata':ds.nodata}
except Exception as e:report['D_DEM_error']=str(e)
for p in (ROOT/'F题exp'/'real_attachments').rglob('*.csv'):
    with p.open(encoding='utf-8-sig',newline='') as f:
        rr=csv.DictReader(f); rows=list(rr)
    report.setdefault('F_csv',{})[str(p.relative_to(ROOT/'F题exp'))]={'rows':len(rows),'columns':rr.fieldnames,'sample':rows[:2]}
zp=next((ROOT/'E题exp').glob('*.zip'))
with zipfile.ZipFile(zp) as z:
    report['E_nonvideo']=[{'name':x.filename,'bytes':x.file_size} for x in z.infolist() if not x.is_dir() and not x.filename.endswith('.mp4')]
    report['E_labels']={n:excel_info(io.BytesIO(z.read(n))) for n in z.namelist() if n.endswith('.xlsx')}
(OUT/'data_inspection.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k not in ['F_csv','E_labels','D_excel','B_excel','E_nonvideo','A_graphs']},ensure_ascii=False,indent=2))
print('A graph count',len(report['A_graphs']))
print('A graph node ranges',report['A_graphs'][:3],report['A_graphs'][-3:])
print('Inspection saved',OUT/'data_inspection.json')
